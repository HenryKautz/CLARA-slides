"""Build CLARA-slides.pptx: plain 16:9 deck, native (editable) equations from LaTeX."""
import copy, re, subprocess, sys, tempfile, zipfile, os
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

OUT = sys.argv[1]
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
      "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
      "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006"}

prs = Presentation()
scale = 13.333 / 10
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
# layouts before the master: a layout placeholder inherits its position from
# the master, so scaling the master first would scale those twice
for owner in [*prs.slide_layouts, prs.slide_master]:
    for sh in owner.placeholders:
        l, t, w, h = sh.left, sh.top, sh.width, sh.height
        sh.left, sh.top, sh.width, sh.height = int(l * scale), t, int(w * scale), h
LAYOUT = prs.slide_layouts[1]  # Title and Content


def runs(par, text, size):
    """Add text to a paragraph; **bold** and *italic* markup."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok:
            continue
        r = par.add_run()
        if tok.startswith("**"):
            r.text, r.font.bold = tok[2:-2], True
        elif tok.startswith("*"):
            r.text, r.font.italic = tok[1:-1], True
        else:
            r.text = tok
        r.font.size = Pt(size)


def bullets(shape, items, size=24, y=None, h=None):
    """Fill a placeholder (or textbox) with (level, text) items."""
    if y is not None:
        shape.left, shape.top, shape.width, shape.height = Inches(0.8), Inches(y), Inches(11.7), Inches(h)
    tf = shape.text_frame
    tf.word_wrap = True
    for i, (lvl, text) in enumerate(items):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.level = lvl
        par.space_before = Pt(6 if lvl == 0 else 2)
        runs(par, text, size if lvl == 0 else size - 4)


def body_box(slide, items, y, h, size=24):
    """A second bulleted region: a copy of the layout's body placeholder."""
    ph = slide.placeholders[1]
    el = copy.deepcopy(ph._element)
    ph._element.addnext(el)
    shape = slide.shapes[-1] if False else None
    # find the shape object wrapping the new element
    for s in slide.shapes:
        if s._element is el:
            shape = s
    for p in list(shape.text_frame._txBody.findall("a:p", NS)):
        shape.text_frame._txBody.remove(p)
    etree.SubElement(shape.text_frame._txBody, f"{{{NS['a']}}}p")
    el.find(".//p:cNvPr", NS).set("id", str(900 + len(slide.shapes)))
    el.find(".//p:cNvPr", NS).set("name", f"Body {len(slide.shapes)}")
    bullets(shape, items, size, y, h)
    return shape


def omml(latex):
    """LaTeX -> the <a14:m> element pandoc writes for display math in pptx."""
    with tempfile.TemporaryDirectory() as d:
        src, out = os.path.join(d, "e.md"), os.path.join(d, "e.pptx")
        open(src, "w").write("## x\n\n$$" + latex + "$$\n")
        subprocess.run(["pandoc", src, "--slide-level=2", "-o", out], check=True)
        xml = zipfile.ZipFile(out).read("ppt/slides/slide1.xml")
    return etree.fromstring(xml).find(".//a14:m", NS)


def equation(slide, latex, y, h, size=24):
    """Equation box: native PowerPoint equation, LaTeX text as the fallback."""
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.7), Inches(h))
    tb.name = "Equation"
    tb.text_frame.word_wrap = True
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    par = tb.text_frame.paragraphs[0]
    par.alignment = PP_ALIGN.CENTER
    runs(par, latex, size)
    sp = tb._element
    choice_sp = copy.deepcopy(sp)
    p = choice_sp.find(".//a:p", NS)
    for r in p.findall("a:r", NS):
        p.remove(r)
    m = omml(latex)
    for mr in m.iter(f"{{{NS['m']}}}r"):  # font size for every math run
        rpr = etree.Element(f"{{{NS['a']}}}rPr", sz=str(size * 100))
        etree.SubElement(rpr, f"{{{NS['a']}}}latin", typeface="Cambria Math")
        mt = mr.find("m:t", NS)
        mt.addprevious(rpr)
    p.append(m)
    ac = etree.Element(f"{{{NS['mc']}}}AlternateContent", nsmap={"mc": NS["mc"]})
    ch = etree.SubElement(ac, f"{{{NS['mc']}}}Choice", Requires="a14", nsmap={"a14": NS["a14"]})
    ch.append(choice_sp)
    fb = etree.SubElement(ac, f"{{{NS['mc']}}}Fallback")
    sp.addprevious(ac)
    fb.append(sp)


def table(slide, rows, y, h, widths, size=18):
    nr, nc = len(rows), len(rows[0])
    t = slide.shapes.add_table(nr, nc, Inches(0.8), Inches(y), Inches(sum(widths)), Inches(h)).table
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = t.cell(i, j)
            cell.text_frame.word_wrap = True
            par = cell.text_frame.paragraphs[0]
            par.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            runs(par, text, size)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def text(slide, s, y, h=0.6, size=20):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.7), Inches(h))
    tb.text_frame.word_wrap = True
    runs(tb.text_frame.paragraphs[0], s, size)


def new(title):
    s = prs.slides.add_slide(LAYOUT)
    s.shapes.title.text = title
    return s


# 1 -------------------------------------------------------------------------
s = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide
s.shapes.title.text = "CLARA Monthly Tagup"
sub = s.placeholders[1].text_frame
for i, line in enumerate(["October 20, 2026", "",
                          "Henry Kautz, University of Virginia",
                          "Steven Holtzen, Northeastern University",
                          "Zhe Zeng, University of Virginia"]):
    par = sub.paragraphs[0] if i == 0 else sub.add_paragraph()
    runs(par, line, 20)

s = new("Planning and Plan Recognition")
bullets(s.placeholders[1], [
    (0, "**Planning:** from an initial state, find actions that achieve a goal"),
    (0, "**PDDL:** standard language for planning domains and problems"),
    (1, "objects, predicates, actions with preconditions and effects"),
    (0, "**Action costs, goal utilities:** a plan's value is the utility of the goals it reaches minus the cost of its actions"),
    (0, "**Disjunctive goals:** the agent pursues *one* of several possible goals"),
    (0, "**Plan recognition** runs planning in reverse: observe some actions, infer which goal the agent is pursuing"),
], y=1.6, h=5.4, size=28)

# 2 -------------------------------------------------------------------------
s = new("Domain: Logistics")
bullets(s.placeholders[1], [
    (0, "Packages, trucks, and airplanes"),
    (1, "trucks drive between places on roads"),
    (1, "airplanes fly between airports"),
    (1, "packages are loaded, unloaded, and transferred between vehicles"),
    (0, "Driving and flying have costs; loading is free"),
    (0, "Finding a minimum-cost plan is NP-complete"),
    (1, "one truck delivering packages to *n* cities, with a cost on each road, is the traveling salesman problem"),
], y=1.6, h=5.4, size=28)

# 3 -------------------------------------------------------------------------
s = new("From Planning to Plan Recognition")
bullets(s.placeholders[1], [
    (0, "How do we get a probability distribution from a planning problem plus evidence?"),
    (0, "Treat every plan τ up to horizon *H* as possible, with probability exponential in its net utility:"),
], y=1.5, h=1.6, size=22)
equation(s, r"P(\tau) = \frac{e^{U(\tau)}}{Z}, \qquad U(\tau) = \sum_{g \in \mathrm{goals}(\tau)} u_g \;-\; \sum_{a \in \tau} c_a", 3.1, 1.2)
body_box(s, [(0, "Probability of goal *g* given observations *O*  (τ ⊨ O ∧ g: plan τ fits *O* and reaches *g*):")], 4.4, 0.6, size=22)
equation(s, r"P(g \mid O) = \frac{\sum_{\tau \models O \wedge g} e^{U(\tau)}}{\sum_{\tau \models O} e^{U(\tau)}}", 5.1, 1.5)

# 4 -------------------------------------------------------------------------
s = new("Example")
for k, lab in enumerate(["g1", "start", "", "g2"]):
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.2 + 2.3 * k), Inches(1.6), Inches(2.1), Inches(0.9))
    b.name = f"Cell {k + 1}"
    b.fill.solid()
    b.fill.fore_color.rgb = RGBColor(0xDD, 0xE8, 0xF5) if lab.startswith("g") else RGBColor(0xFF, 0xFF, 0xFF)
    b.line.color.rgb = RGBColor(0x40, 0x40, 0x40)
    b.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    par = b.text_frame.paragraphs[0]
    par.alignment = PP_ALIGN.CENTER
    if lab:
        runs(par, lab, 24)
        par.runs[0].font.color.rgb = RGBColor(0, 0, 0)
text(s, "Each move costs 1 · each goal is worth +4 · horizon 3 (waiting is free)", 2.8)
table(s, [
    ["Plans ending at", "No evidence", "After a move right"],
    ["g1", "**0.72**  (3 plans of U = 3, 2 of U = 1)", "**0.15**  (1 plan of U = 1)"],
    ["g2", "**0.24**  (3 plans of U = 2)", "**0.81**  (2 plans of U = 2)"],
    ["neither", "0.03  (13 plans)", "0.04  (5 plans)"],
], 3.6, 2.0, [2.5, 4.6, 4.6])
text(s, "g1 is still possible after moving right, but much less likely", 6.0)
s.placeholders[1]._element.getparent().remove(s.placeholders[1]._element)

# 5 -------------------------------------------------------------------------
s = new("Baseline: Translation to FiFO")
bullets(s.placeholders[1], [
    (0, "**FiFO:** finite-domain first-order logic, compiled to propositional CNF"),
    (0, "**Weighted logic (Markov logic):** each weighted formula adds its cost when true"),
], y=1.5, h=1.3, size=22)
equation(s, r"P(x) \propto \exp\Big(-\sum_{i \,:\, x \models f_i} w_i\Big)", 2.8, 1.0)
body_box(s, [
    (0, "**SATPLAN:** a PDDL problem at horizon *H* becomes CNF over atoms *holds(f, t)* and *occurs(a, t)*"),
    (1, "each satisfying assignment is a plan; action costs become weights"),
    (0, "Plan recognition is then a ratio of weighted model counts:"),
], 3.8, 1.9, size=22)
equation(s, r"P(g \mid O) = \frac{\mathrm{WMC}(T \wedge O \wedge g)}{\mathrm{WMC}(T \wedge O)}", 5.8, 1.1)

# 6 -------------------------------------------------------------------------
s = new("Model Counting")
bullets(s.placeholders[1], [
    (0, "**Exact: SharpSAT-TD**"),
    (1, "DPLL search with component caching, guided by a tree decomposition"),
    (1, "exact answers; #P-hard in general"),
    (0, "**Approximate: Monte Carlo (MC-SAT)**"),
    (1, "Markov chain whose steps sample satisfying assignments with a SAT-based sampler"),
    (1, "all marginals from one run; can mix slowly"),
], y=1.6, h=5.4, size=28)

# 7 -------------------------------------------------------------------------
s = new("Max-Model Approximation")
bullets(s.placeholders[1], [
    (0, "Replace each sum over plans by its largest term: the cheapest plan (Ramírez & Geffner 2010)"),
    (0, "For each goal *g*, two MaxSAT calls:"),
    (1, "c(g, O): cheapest plan to *g* consistent with the observations"),
    (1, "c(g, ¬O): cheapest plan to *g* that is not"),
], y=1.5, h=2.4, size=22)
equation(s, r"P(g \mid O) \propto \sigma\big(\beta\,(c(g,\neg O) - c(g, O))\big)", 4.0, 0.9)
body_box(s, [
    (0, "No counting, only optimization: 2*n* MaxSAT calls for *n* goals"),
    (0, "Ignores *how many* plans there are, so goals with equally cheap plans tie"),
], 5.0, 1.6, size=22)

# 8 -------------------------------------------------------------------------
s = new("Results on Simple Grids")
table(s, [
    ["Grid", "Goals", "CNF vars", "Max-term\nP / time", "SharpSAT-TD\nP / time", "MC-SAT\nP / time"],
    ["3×3", "2", "141", "0.50 / 12 s", "0.50 / 16 s", "0.51 / 14 s"],
    ["4×4", "3", "400", "0.33 / 16 s", "0.37 / 20 s", "0.36 / 33 s"],
    ["5×5", "4", "812", "0.26 / 35 s", "0.27 / 35 s", "0.25 / 210 s"],
    ["6×6", "5", "1,557", "0.20 / 30 s", "0.21 / 38 s", "0.19 / 349 s"],
    ["7×7", "6", "2,436", "0.18 / 40 s", "0.20 / 105 s †", "0.19 / 737 s †"],
    ["8×8", "7", "4,024", "0.15 / 61 s", "none finished", "none finished"],
], 1.5, 3.6, [1.3, 1.2, 1.6, 2.5, 2.5, 2.5], size=16)
body_box(s, [
    (0, "P = mean probability given to the true goal (chance = 1 / goals); time = mean per run"),
    (0, "Counting does not scale: † on 7×7, 36 SharpSAT-TD and 2 MC-SAT runs of 144 hit the 20-minute limit; none finished on 8×8"),
    (0, "Max-term always finishes but ties every goal: early moves on an open grid say little about the goal"),
], 5.25, 1.9, size=18)

# 9 -------------------------------------------------------------------------
s = new("Knowledge Compilation")
bullets(s.placeholders[1], [
    (0, "Compile the problem once to a d-DNNF circuit (d4), with no evidence"),
    (0, "Then marginals given *any* observed moves come from one pass over the circuit"),
], y=1.5, h=1.2, size=22)
table(s, [
    ["Grid", "Horizon", "CNF vars", "Compile time", "Circuit nodes", "Query, 2 obs."],
    ["3×3", "5", "141", "0.04 s", "44", "0.001 s"],
    ["5×5", "8.5", "812", "0.09 s", "1,108", "0.007 s"],
    ["6×6", "10.75", "1,557", "1.4 s", "21,140", "0.017 s"],
    ["7×7", "12", "2,436", "6.8 s", "82,748", "0.035 s"],
    ["8×8", "14.75", "4,024", "104 s", "785,660", "0.24 s"],
    ["9×9", "16", "5,616", "418 s", "2,525,681", "0.45 s"],
], 2.8, 3.3, [1.3, 1.5, 1.8, 2.3, 2.4, 2.4], size=16)
text(s, "10×10 did not compile within 20 minutes", 6.3, size=20)

# 10 ------------------------------------------------------------------------
# CLARA/TwoPackage/results/{runs,compile-d4,query-evidence}.jsonl, last attempt
# of each run (the 20-minute rerun where there was one); 4 instances per size
s = new("Results on Two-Package Grids")
text(s, "One truck, two packages, two possible destinations each; the agent delivers one "
        "package: 4 goals (chance = 0.25), 4 instances per grid, 1 … N−2 observed actions", 1.4, h=0.9, size=20)
table(s, [
    ["Grid", "CNF vars", "Max-term\nP / time", "SharpSAT-TD\nP / time", "MC-SAT\nP / time", "d4 compile\ndone · time"],
    ["3×3", "752", "0.40 / 24 s", "0.46 / 28 s", "0.37 / 85 s", "4 of 4 · 4 s"],
    ["4×4", "1,893", "0.35 / 35 s", "0.39 / 52 s", "0.31 / 454 s", "4 of 4 · 5 s"],
    ["5×5", "3,147", "0.39 / 56 s", "0.41 / 67 s", "0.32 / 579 s †", "4 of 4 · 9 s"],
    ["6×6", "5,797", "0.39 / 103 s", "none finished", "not run", "4 of 4 · 105 s"],
    ["7×7", "9,925", "0.39 / 355 s", "not run", "not run", "2 of 4 · 207 s"],
], 2.25, 2.8, [1.1, 1.5, 2.25, 2.25, 2.3, 2.3], size=16)
body_box(s, [
    (0, "P = mean probability given to the true goal over finished runs; time = mean per run; 20-minute limit; a method that timed out was not run on larger grids"),
    (0, "Counting stops early: SharpSAT-TD finished none of 8 runs on 6×6; † MC-SAT finished 4 of 16 on 5×5"),
    (0, "Max-term finishes everything, within 0.06 of exact counting, but its top goal is a tie in 216 of 240 runs"),
    (0, "d4 compiles through 6×6; each query then takes ≤ 0.26 s and matches SharpSAT-TD to 2·10⁻¹⁵"),
], 5.55, 1.85, size=17)

# 11 ------------------------------------------------------------------------
s = new("Test Problem Archive")
bullets(s.placeholders[1], [
    (0, "**GitHub:** github.com/HenryKautz/CLARA"),
    (1, "tools: github.com/HenryKautz/FiFO"),
    (0, "**Problem generator:** ppgen makes problems, planner finds plans, evgen turns plans into observations"),
    (0, "**Simple grid instances:** N×N grids, N = 3 … 10, one truck, several possible destinations"),
    (0, "**Two-package instances:** N×N grids, N = 3 … 7, two packages with two possible destinations each"),
    (0, "**Baseline results:** the tables on the previous three slides"),
], y=1.6, h=5.4, size=28)

prs.save(OUT)
